"""Strict input parsing; Excel dimensions are not trusted."""
import datetime as dt
import numbers
import re
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook


def read_table(source, kind, decimal='.'):
    name = getattr(source, 'name', str(source))
    if name.lower().endswith('.csv'):
        df = pd.read_csv(source, sep=None, engine='python', dtype=str, keep_default_na=False)
    else:
        book = load_workbook(source, read_only=True, data_only=True)
        sheet = book['Pagos'] if kind == 'payments' and 'Pagos' in book.sheetnames else book.worksheets[0]
        sheet.reset_dimensions()
        rows = list(sheet.values)
        book.close()
        if not rows:
            raise ValueError('Archivo vacío')
        df = pd.DataFrame(rows[1:], columns=rows[0])
    df.columns = [str(c).strip() for c in df.columns]
    if df.columns.duplicated().any():
        raise ValueError('Encabezados duplicados')
    required = ['Ticker', 'Fecha', 'Renta', 'Amortización', 'Total'] if kind == 'payments' else ['Ticker', 'Market Value', 'Nominal', 'Moneda', 'Emisor']
    missing = set(required) - set(df.columns)
    if missing:
        raise ValueError(f'Faltan columnas: {sorted(missing)}')
    if df.empty:
        raise ValueError('No hay filas de datos')
    issues = []
    def issue(level, row, field, message):
        issues.append(dict(Nivel=level, Archivo=kind, Fila=row + 2, Campo=field, Mensaje=message))
    for col in ['Ticker'] + (['Moneda', 'Emisor'] if kind == 'positions' else []):
        for i, v in df[col].items():
            if pd.isna(v) or not str(v).strip():
                issue('error' if col != 'Emisor' else 'warning', i, col, 'Valor vacío')
                df.at[i, col] = ''
            else:
                df.at[i, col] = str(v).strip()
                if str(v) != str(v).strip():
                    issue('warning', i, col, 'Se eliminaron espacios externos')
    nums = ['Renta', 'Amortización', 'Total'] if kind == 'payments' else ['Market Value', 'Nominal']
    for col in nums:
        for i, v in df[col].items():
            try:
                if pd.isna(v) or str(v).strip() == '':
                    if col == 'Nominal':
                        issue('warning', i, col, 'Nominal faltante; no se usa para escalar flujos')
                        df.at[i, col] = float('nan')
                        continue
                    raise ValueError('Número vacío')
                if isinstance(v, numbers.Real) and not isinstance(v, bool):
                    n = float(v)
                else:
                    text = str(v).strip()
                    pattern = r'[+-]?\d+(?:' + re.escape(decimal) + r'\d+)?'
                    if not re.fullmatch(pattern, text):
                        raise ValueError('Separadores ambiguos: use números sin miles y el decimal elegido')
                    n = float(text.replace(decimal, '.'))
                if not pd.notna(n) or abs(n) == float('inf'):
                    raise ValueError('Número no finito')
                df.at[i, col] = n
                if n < 0 or (col == 'Market Value' and n == 0):
                    issue('error', i, col, 'Negativo inesperado o Market Value no positivo')
            except (ValueError, TypeError) as e:
                issue('error', i, col, str(e)); df.at[i, col] = float('nan')
        df[col] = pd.to_numeric(df[col], errors='coerce')
    if kind == 'payments':
        for i, v in df.Fecha.items():
            try:
                if isinstance(v, (dt.datetime, dt.date, pd.Timestamp)):
                    date = pd.Timestamp(v)
                    if date != date.normalize() or date.tzinfo:
                        raise ValueError('Fecha con hora o zona horaria')
                elif isinstance(v, str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}', v):
                    date = pd.Timestamp(dt.date.fromisoformat(v))
                else:
                    raise ValueError('Use fecha Excel o texto ISO AAAA-MM-DD')
                df.at[i, 'Fecha'] = date
            except (ValueError, TypeError):
                issue('error', i, 'Fecha', 'Fecha inválida o ambigua'); df.at[i, 'Fecha'] = pd.NaT
        df['Fecha'] = pd.to_datetime(df.Fecha)
        df['Flujo Total'] = df.Renta + df['Amortización']
        for i in df.index[(df.Total - df['Flujo Total']).abs() > 0.005]:
            issue('warning', i, 'Total', f"Total informado {df.at[i, 'Total']} vs recalculado {df.at[i, 'Flujo Total']}")
        dup = df.duplicated(['Ticker', 'Fecha'], keep=False)
    else:
        dup = df.duplicated('Ticker', keep=False)
    for i in df.index[dup]:
        issue('error', i, 'Ticker', 'Duplicado inesperado; consolidar explícitamente en el origen')
    return df, issues
